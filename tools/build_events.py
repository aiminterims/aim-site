# -*- coding: utf-8 -*-
"""Читает data/events.xlsx и обновляет блок мероприятий в index.html.
Берутся только строки, где в колонке I стоит x. Используются колонки:
C - дата, D - город, F (или E, если F пуста) - описание,
J (необязательно) - ссылка на страницу записи, например /breakfast/: тогда кнопка «Записаться»."""
import re, sys, html, datetime
from openpyxl import load_workbook

X_COL, DATE_COL, CITY_COL, TEXT_COLS = 'I', 'C', 'D', ('F', 'E')
MONTHS = ['января','февраля','марта','апреля','мая','июня','июля','августа','сентября','октября','ноября','декабря']
CONTACT = 'https://t.me/AndreyAFomin'

def fmt_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return f'{v.day} {MONTHS[v.month-1]}'
    s = str(v or '').strip()
    m = re.match(r'(\d{1,2})\s+([А-Яа-яЁё]+)', s)
    if m:
        d, mon = m.group(1), m.group(2).lower()
        for full in MONTHS:
            if full.startswith(mon[:3]): return f'{d} {full}'
    return s

def city(v):
    s = str(v or '').strip()
    return 'Санкт-Петербург' if s.lower() in ('с-петербург', 'спб', 'санкт-петербург') else s

def sort_key(row):
    m = re.match(r'(\d+)\s+(\S+)', row['date'])
    if not m: return (99, 99)
    return (next((i for i, n in enumerate(MONTHS) if n == m.group(2)), 99), int(m.group(1)))

def main(xlsx='data/events.xlsx', page='index.html'):
    ws = load_workbook(xlsx, data_only=True).active
    rows = []
    for r in range(1, ws.max_row + 1):
        if str(ws[f'{X_COL}{r}'].value or '').strip().lower() != 'x': continue
        text = next((str(ws[f'{c}{r}'].value).strip() for c in TEXT_COLS if ws[f'{c}{r}'].value), '')
        if not text: continue
        rows.append(dict(date=fmt_date(ws[f'{DATE_COL}{r}'].value), city=city(ws[f'{CITY_COL}{r}'].value), text=text, link=str(ws[f'J{r}'].value or '').strip()))
    rows.sort(key=sort_key)
    cards = []
    for e in rows:
        href, label = (e['link'], 'Записаться') if e['link'] else (CONTACT, 'Узнать подробнее')
        cards.append(f'''      <div class="card ev reveal">
        <div class="when"><span class="pill a">{html.escape(e['date'])}</span><span class="pill">{html.escape(e['city'])}</span></div>
        <h3>{html.escape(e['text'])}</h3>
        <a class="btn" href="{html.escape(href)}"{'' if e['link'] else ' target="_blank" rel="noopener"'}>{label}</a>
      </div>''')
    if not cards:
        cards = ['      <div class="card ev reveal"><h3>Ближайшие мероприятия скоро появятся</h3><a class="btn" href="%s" target="_blank" rel="noopener">Следить в Telegram</a></div>' % CONTACT]
    block = '<!--EVENTS-->\n' + '\n'.join(cards) + '\n      <!--/EVENTS-->'
    s = open(page, encoding='utf-8').read()
    s2, n = re.subn(r'<!--EVENTS-->.*?<!--/EVENTS-->', lambda m: block, s, flags=re.S)
    if n != 1: sys.exit('Маркеры EVENTS не найдены')
    open(page, 'w', encoding='utf-8').write(s2)
    print(f'Мероприятий на сайте: {len(rows)}')

if __name__ == '__main__':
    main(*sys.argv[1:])
